import os
import json
import re
from datetime import datetime
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ---- Dátum és fájl elérési útvonalak ----
TODAY = datetime.now().strftime("%Y_%m_%d")
BASE_DIR = os.path.dirname(__file__)
INPUT_FILE = os.path.join(BASE_DIR, "..", "output", f"{TODAY}_vatera.json")
OUTPUT_FILE = os.path.join(BASE_DIR, "..", "output", f"{TODAY}_vatera_export.xlsx")

# ---- JSON beolvasás ----
with open(INPUT_FILE, "r", encoding="utf-8") as f:
    data = json.load(f)

# ---- DataFrame konvertálás ----
df = pd.DataFrame(data)

# ---- Illegális karakterek eltávolítása minden szöveges oszlopból ----
def clean_text(val):
    if isinstance(val, str):
        return re.sub(r"[\x00-\x1F\x7F]", "", val)
    return val

df = df.applymap(clean_text)

# ---- Excel írás ----
with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="Vatera Adatok")
    ws = writer.sheets["Vatera Adatok"]

    # Táblázat kinézet és stílus
    num_cols = len(df.columns)
    last_col = get_column_letter(num_cols)
    last_row = len(df) + 1
    table_range = f"A1:{last_col}{last_row}"
    table = Table(displayName="VateraTable", ref=table_range)

    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table.tableStyleInfo = style
    ws.add_table(table)

print(f"✅ Excel fájl létrehozva: {OUTPUT_FILE}")